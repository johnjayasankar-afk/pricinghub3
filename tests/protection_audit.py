"""Static audit: every light-blue input cell is unlocked and every formula cell is locked, so Review -> Protect Sheet works as documented."""
import os, sys
from openpyxl import load_workbook
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INPUT = "00EAF3FF"
DEFAULT_FORMULA_INPUTS = {("Schedule", "B2"), ("Timesheet", "Q3")}  # unlocked inputs that ship with a default formula (this Monday / last Sunday) the user may overwrite
bad = []; n_in = n_fx = 0
for fn in ("Contractor-Job-Costing-SAMPLE.xlsx", "Contractor-Job-Costing-BLANK.xlsx"):
    wb = load_workbook(os.path.join(ROOT, "product", fn))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                is_input = c.fill is not None and c.fill.fill_type == "solid" and str(c.fill.start_color.rgb) == INPUT
                is_fx = isinstance(c.value, str) and c.value.startswith("=")
                if is_input:
                    n_in += 1
                    if c.protection.locked: bad.append((fn, ws.title, c.coordinate, "input cell is locked"))
                    if is_fx and (ws.title, c.coordinate) not in DEFAULT_FORMULA_INPUTS: bad.append((fn, ws.title, c.coordinate, "input-styled cell holds a formula"))
                elif is_fx:
                    n_fx += 1
                    if not c.protection.locked: bad.append((fn, ws.title, c.coordinate, "formula cell is unlocked"))
print(f"input cells: {n_in}  formula cells: {n_fx}  problems: {len(bad)}")
for b in bad[:20]: print(" ", b)
sys.exit(1 if bad else 0)
