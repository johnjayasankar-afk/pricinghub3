"""Static Google Sheets / Excel 2016 compatibility audit: every formula, conditional-format rule and data validation in the workbook."""
import re, sys, os
from openpyxl import load_workbook
from collections import Counter
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHEETS_OK = {"SUM","SUMIFS","SUMIF","COUNTIF","COUNTIFS","COUNTA","COUNT","VLOOKUP","INDEX","MATCH","IFERROR","IF","AND","OR","NOT","MAX","MIN",
             "SUMPRODUCT","ISERROR","ISNUMBER","DATE","EOMONTH","TODAY","WORKDAY","WEEKDAY","DAY","MONTH","YEAR","TEXT","LEFT","ROUND","ABS","LEN","N"}
EXCEL2016_OK = SHEETS_OK  # all of the above exist in Excel 2010+
EXCEL_ONLY_MARKERS = [r"_xlfn\.", r"(?<![A-Z])XLOOKUP\(", r"(?<![A-Z])FILTER\(", r"(?<![A-Z])UNIQUE\(", r"(?<![A-Z])LET\(", r"(?<![A-Z])TEXTJOIN\(", r"(?<![A-Z])IFS\(", r"(?<![A-Z])MAXIFS\(", r"(?<![A-Z])MINIFS\(", r"(?<![A-Z])SWITCH\(", r"@", r"\[#", r"\[@"]
wb = load_workbook(os.path.join(ROOT, "product", "Contractor-Job-Costing-SAMPLE.xlsx"))
funcs = Counter(); n_formulas = 0; problems = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                n_formulas += 1
                for m in EXCEL_ONLY_MARKERS:
                    if re.search(m, v): problems.append((ws.title, c.coordinate, m))
                for f in re.findall(r"([A-Z][A-Z0-9\.]*)\(", v):
                    funcs[f] += 1
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            for f in (rule.formula or []):
                for fn in re.findall(r"([A-Z][A-Z0-9\.]*)\(", str(f)): funcs["CF:" + fn] += 1
                for m in EXCEL_ONLY_MARKERS:
                    if re.search(m, str(f)): problems.append((ws.title, str(cf.sqref), m))
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list": problems.append((ws.title, str(dv.sqref), f"non-list validation {dv.type}"))
unknown = {f for f in funcs if not f.startswith("CF:") and f not in SHEETS_OK}
unknown_cf = {f[3:] for f in funcs if f.startswith("CF:") and f[3:] not in SHEETS_OK}
print(f"formulas scanned: {n_formulas}")
print("functions used:", ", ".join(sorted(f for f in funcs if not f.startswith('CF:'))))
print("functions in conditional formats:", ", ".join(sorted(f[3:] for f in funcs if f.startswith('CF:'))))
print("not in the Google Sheets allowlist:", sorted(unknown | unknown_cf) or "none")
print("Excel-only constructs found:", problems or "none")
print("defined names:", list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else [d.name for d in wb.defined_names.definedName])
sys.exit(1 if (unknown or unknown_cf or problems) else 0)
