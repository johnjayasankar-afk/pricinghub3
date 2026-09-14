"""Adversarial data test: inject bad rows into the SAMPLE workbook, recalculate in Excel, and prove every Health Check line,
the Check columns and the Dashboard alert respond. Complements test_workbook.py (which proves the numbers on clean data)."""
import os, sys, json, subprocess, datetime as dt, shutil
from openpyxl import load_workbook
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "product", "Contractor-Job-Costing-SAMPLE.xlsx")
OUT = os.path.join(ROOT, "tests", "_adversarial.xlsx")
M = json.load(open(os.path.join(ROOT, "product", "layout.json"))); C = M["C"]; R1 = M["R1"]; hc = M["hc"]; VR = M["VEND_R1"]
wb = load_workbook(SRC)
def put(sheet, row, **cols):
    ws = wb[sheet]; cmap = C[sheet] if sheet in C else None
    for h, v in cols.items():
        col = cmap[h] if cmap else h
        ws[f"{col}{row}"] = v
D = dt.date
# ---- one bad row per input sheet (rows sit just below the sample data) ----
put("Jobs", 9, **{"Job ID": "J-1009", "Client": "Nobody Inc", "Status": None, "Start date": D(2026, 9, 1)})                # unknown client
put("Jobs", 10, **{"Job ID": "J-1001", "Client": "Kim & Patel", "Status": "Lead"})                                            # duplicate ID
put("Clients", 8, **{"Client": "Alvarez residence", "Contact person": "Dup"})                                                # duplicate client
wb["Vendors & Crew"]["A9"] = "M. Ortiz"                                                                                        # duplicate crew, no rate
wb["Vendors & Crew"][f"A{VR + 11}"] = "Volt Electric"; wb["Vendors & Crew"][f"B{VR + 11}"] = "Subcontractor"                  # duplicate vendor
put("Estimate", 37, **{"Job ID": "J-9999", "Cost Code": "05", "Description": "Ghost job line", "Qty": 1, "Unit Cost": 100})   # unknown job
put("Schedule", 21, **{"Job ID": "J-9999", "Task": "Ghost task", "Duration (workdays)": 2})                                    # unknown job, no start
put("Timesheet", 28, **{"Date": D(2026, 8, 3), "Worker": "Ghost Worker", "Job ID": "J-1001", "Cost Code": "05", "Hours": 8})   # unknown worker
put("Job Costs", 24, **{"Job ID": "J-1001", "Cost Code": "ZZ", "Vendor": "Paint Pro", "Description": "Bad code", "Qty or Hours": 1, "Unit Cost / Rate": 50, "Paid?": "Yes"})  # bad code, no date
put("Job Costs", 25, **{"Date": D(2026, 6, 20), "Job ID": "J-1002", "Cost Code": "06", "Vendor": "Builders Supply", "Description": "Budget buster", "Qty or Hours": 1, "Unit Cost / Rate": 1000, "Paid?": "Yes"})  # pushes J-1002 over budget
put("Mileage", 8, **{"Date": D(2026, 8, 1), "Job ID": "J-9999", "From": "A", "To": "B"})                                     # unknown job, no miles
put("Daily Log", 9, **{"Date": D(2026, 8, 1), "Job ID": "J-9999", "Weather": "Sunny"})                                          # unknown job
put("Change Orders", 9, **{"CO #": 1, "Job ID": "J-1001", "Date": D(2026, 8, 20), "Description": "Dup number", "Status": "Pending", "Added Cost (budget)": 10})           # duplicate CO #
put("Change Orders", 10, **{"CO #": 7, "Job ID": "J-1001", "Date": D(2026, 8, 21), "Description": "No client date", "Status": "Approved", "Added Cost (budget)": 10})   # approved w/o client date
put("Invoices", 10, **{"Invoice #": "INV-2026-014", "Job ID": "J-1001", "Invoice Date": D(2026, 8, 30), "Due Date": D(2026, 9, 13), "Amount": 100, "Description": "Duplicate number"})
put("Invoices", 11, **{"Invoice #": "INV-2026-030", "Job ID": "J-1001", "Invoice Date": D(2026, 6, 1), "Due Date": D(2026, 6, 15), "Amount": 500, "Description": "Old and overdue"})
put("Payments", 9, **{"Date": D(2026, 9, 1), "Job ID": "J-1001", "Invoice #": "INV-2026-023", "Amount": 9000, "Method": "Check"})   # overpays an $8,000 invoice
put("Payments", 10, **{"Date": D(2026, 9, 1), "Job ID": "J-9999", "Invoice #": "INV-0000", "Amount": 5, "Method": "Cash"})          # unknown job and invoice
wb.save(OUT)
cells = {"hc_total": hc["total_issues"], "hc_status": f"D{hc['total_issues_row']}", "diag_total": hc["diag_total"]}
for i, name in enumerate(["jobs", "clients", "crew", "vendors", "estimate", "schedule", "timesheet", "jobcosts", "mileage", "dailylog", "cos", "invoices", "payments"]): cells[f"hc_{name}"] = f"C{6 + i}"
for i, name in enumerate(["over_budget", "overdue", "late", "expired", "expiring", "co_nodate", "unpaid", "overpaid"]): cells[f"att_{name}"] = f"C{hc['info_first'] + i}"
checks = {"dash_issues": ("Dashboard", f"K{M['dash']['alert_row']}"), "dash_overdue": ("Dashboard", f"B{M['dash']['alert_row']}"), "chk_jobs9": ("Jobs", f"{C['Jobs']['Check']}9"), "chk_jobs10": ("Jobs", f"{C['Jobs']['Check']}10"), "chk_clients8": ("Clients", f"{C['Clients']['Check']}8"),
          "chk_crew9": ("Vendors & Crew", f"{M['CC']['Check']}9"), "chk_vend": ("Vendors & Crew", f"{M['VC']['Check']}{VR + 11}"), "chk_est37": ("Estimate", f"{C['Estimate']['Check']}37"),
          "chk_sch21": ("Schedule", f"{C['Schedule']['Check']}21"), "chk_ts28": ("Timesheet", f"{C['Timesheet']['Check']}28"), "rate_ts28": ("Timesheet", f"{C['Timesheet']['Rate Used']}28"),
          "chk_jc24": ("Job Costs", f"{C['Job Costs']['Check']}24"), "chk_mi8": ("Mileage", f"{C['Mileage']['Check']}8"), "chk_dl9": ("Daily Log", f"{C['Daily Log']['Check']}9"),
          "chk_co9": ("Change Orders", f"{C['Change Orders']['Check']}9"), "chk_co10": ("Change Orders", f"{C['Change Orders']['Check']}10"), "chk_inv10": ("Invoices", f"{C['Invoices']['Check']}10"),
          "status_inv11": ("Invoices", f"{C['Invoices']['Status']}11"), "bucket_inv11": ("Invoices", f"{C['Invoices']['Aging Bucket']}11"), "bal_inv7": ("Invoices", f"{C['Invoices']['Balance']}7"),
          "chk_pay10": ("Payments", f"{C['Payments']['Check']}10"), "pct_j1002": ("Jobs", f"{C['Jobs']['% of Budget Used']}6")}
refs = [f"Health Check!{a}" for a in cells.values()] + [f"{s}!{a}" for s, a in checks.values()]
r = subprocess.run(["osascript", os.path.join(ROOT, "tests", "excel_recalc.applescript"), OUT] + refs, capture_output=True, text=True)
if r.returncode: print(r.stderr); sys.exit(1)
got = {}
for line in r.stdout.strip().split("\n"):
    k, v = line.split("\t", 1); got[k] = v
val = {name: got[f"Health Check!{a}"] for name, a in cells.items()}
val.update({name: got[f"{s}!{a}"] for name, (s, a) in checks.items()})
def num(x): return float(x) if x not in ("", None) else 0.0
fails = []
def expect(cond, label, actual):
    if not cond: fails.append(f"{label}: got {actual!r}")
for k in ["hc_jobs", "hc_clients", "hc_crew", "hc_vendors", "hc_estimate", "hc_schedule", "hc_timesheet", "hc_jobcosts", "hc_mileage", "hc_dailylog", "hc_cos", "hc_invoices", "hc_payments"]:
    expect(num(val[k]) >= 1, f"Health Check line {k} should count at least one issue", val[k])
expect(val["hc_status"] == "Review", "Total status should read Review", val["hc_status"])
expect(num(val["hc_total"]) == sum(num(val[k]) for k in val if k.startswith("hc_") and k not in ("hc_total", "hc_status")), "Total equals the sum of lines", val["hc_total"])
expect(num(val["dash_issues"]) == num(val["hc_total"]), "Dashboard data-issues tile mirrors Health Check total", val["dash_issues"])
expect(num(val["diag_total"]) == 0, "No formula errors even with bad data", val["diag_total"])
for k in ["att_over_budget", "att_overdue", "att_late", "att_expired", "att_expiring", "att_co_nodate", "att_unpaid", "att_overpaid"]:
    expect(num(val[k]) >= 1, f"Attention item {k} should fire", val[k])
expect(num(val["dash_overdue"]) == num(val["att_overdue"]), "Dashboard overdue tile mirrors Health Check", val["dash_overdue"])
expect("Client" in val["chk_jobs9"], "Unknown client flagged on Jobs", val["chk_jobs9"]); expect("Duplicate" in val["chk_jobs10"], "Duplicate Job ID flagged", val["chk_jobs10"])
expect("Duplicate" in val["chk_clients8"], "Duplicate client flagged", val["chk_clients8"]); expect(val["chk_crew9"] != "", "Duplicate crew / missing rate flagged", val["chk_crew9"])
expect("Duplicate" in val["chk_vend"], "Duplicate vendor flagged", val["chk_vend"]); expect("Unknown" in val["chk_est37"], "Unknown job on Estimate", val["chk_est37"])
expect(val["chk_sch21"] != "", "Schedule unknown job / missing start", val["chk_sch21"]); expect("orker" in val["chk_ts28"], "Unknown worker on Timesheet", val["chk_ts28"])
expect(num(val["rate_ts28"]) == 0, "Unknown worker gets rate 0 rather than an error", val["rate_ts28"])
expect(val["chk_jc24"] != "", "Bad cost code / missing date on Job Costs", val["chk_jc24"]); expect(val["chk_mi8"] != "", "Mileage unknown job / missing miles", val["chk_mi8"])
expect("Unknown" in val["chk_dl9"], "Daily Log unknown job", val["chk_dl9"]); expect("Duplicate" in val["chk_co9"], "Duplicate CO #", val["chk_co9"])
expect(val["chk_co10"] == "Approved without client date", "Approved CO without client date", val["chk_co10"]); expect("Duplicate" in val["chk_inv10"], "Duplicate invoice #", val["chk_inv10"])
expect(val["status_inv11"] == "Overdue", "Old unpaid invoice is Overdue", val["status_inv11"]); expect(val["bucket_inv11"] == "61-90", "82 days past due lands in 61-90", val["bucket_inv11"])
expect(num(val["bal_inv7"]) < 0, "Overpaid invoice shows negative balance", val["bal_inv7"]); expect(val["chk_pay10"] != "", "Payment with unknown job/invoice flagged", val["chk_pay10"])
expect(num(val["pct_j1002"]) > 1, "J-1002 pushed over budget", val["pct_j1002"])
os.remove(OUT)
if fails:
    print("\n".join("FAIL " + f for f in fails)); print(f"\n{len(fails)} adversarial expectations failed"); sys.exit(1)
print(f"adversarial: all {len(val)} probes behaved as designed; Health Check total = {val['hc_total']}, formula errors = {val['diag_total']}")
