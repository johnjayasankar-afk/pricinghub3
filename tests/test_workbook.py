"""v2 cross-check: recomputes every major output from sample_data.py in pure Python and compares with Microsoft Excel's
recalculated values (via AppleScript). Exit 1 on any mismatch."""
import subprocess, sys, os, json, datetime as dt
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "product"))
import sample_data as SD
from layout import *
XLSX = os.path.join(ROOT, "product", "Contractor-Job-Costing-SAMPLE.xlsx")
M = json.load(open(os.path.join(ROOT, "product", "layout.json")))
TODAY = dt.date.today(); YEAR = 2026
markup = dict(CATS)
from openpyxl import load_workbook
_wb = load_workbook(XLSX); st = _wb["Settings"]
codes = {st.cell(row=r, column=5).value: st.cell(row=r, column=7).value for r in range(CODE_R1, CODE_R2 + 1)}
cont, tax, dep, prog, grace, mil_rate, valid_days, overhead, burden, tax_scope = [st[f"B{r}"].value for r in (9, 10, 11, 12, 13, 15, 16, 17, 18, 19)]
lib = {i[0]: i for i in SD.LIBRARY}; crew = {c[0]: c for c in SD.CREW}
def cat_of(code, default="Other"): return codes.get(code, default) if code else default
est = []
for job, code, item, desc, qty, unit, ucost in SD.ESTIMATE:
    cat = cat_of(code); uc = ucost if ucost is not None else lib[item][3]; cost = qty * uc
    est.append(dict(job=job, code=code, item=item, desc=desc or item, qty=qty, unit=unit or lib[item][2], cost=cost, price=cost * (1 + markup[cat]), cat=cat))
def esum(job, key, **f): return sum(e[key] for e in est if e["job"] == job and all(e[k] == v for k, v in f.items()))
cos = []
for n, job, d, desc, status, mk, cost, sched, appr in SD.CHANGE_ORDERS:
    cos.append(dict(n=n, job=job, status=status, cost=cost, price=cost * (1 + (markup["Other"] if mk is None else mk)), sched=sched))
def co_price(job): return sum(c["price"] for c in cos if c["job"] == job and c["status"] == "Approved")
def co_cost(job): return sum(c["cost"] for c in cos if c["job"] == job and c["status"] == "Approved")
ts = []
for d, w, job, code, hrs, ovr in SD.TIMESHEET:
    rate = ovr if ovr is not None else crew[w][2]; bill = crew[w][3]
    we = d - dt.timedelta(days=d.isoweekday() - 1) + dt.timedelta(days=6)
    ts.append(dict(d=d, w=w, job=job, code=code, hrs=hrs, rate=rate, cost=hrs * rate * (1 + burden), bill=bill, billable=hrs * bill, we=we, cat=cat_of(code, "Labor")))
jc = [dict(d=d, job=job, code=code, cat=cat_of(code), vendor=v, qty=q, amt=q * uc, paid=paid) for d, job, code, v, desc, q, uc, paid, *_ in SD.JOB_COSTS]
mil = [dict(d=d, job=job, code=code, cat=cat_of(code), miles=mi, amt=mi * mil_rate) for d, job, code, *_, mi in SD.MILEAGE]
def actual(job, **f):
    return (sum(x["amt"] for x in jc if x["job"] == job and all(x.get(k) == v for k, v in f.items()))
            + sum(x["cost"] for x in ts if x["job"] == job and all(x.get(k) == v for k, v in f.items()))
            + sum(x["amt"] for x in mil if x["job"] == job and all(x.get(k) == v for k, v in f.items())))
pays = [dict(d=d, job=job, inv=inv, amt=a) for d, job, inv, a, *_ in SD.PAYMENTS]
invs = []
for n, job, d, due, amt, desc in SD.INVOICES:
    paid = sum(p["amt"] for p in pays if p["inv"] == n); bal = amt - paid
    status = "Paid" if bal <= 0 else ("Overdue" if TODAY > due + dt.timedelta(days=grace) else ("Partial" if paid > 0 else "Open"))
    if bal <= 0: bucket = "Paid"
    elif TODAY <= due: bucket = "Current"
    else:
        dd = (TODAY - due).days; bucket = "1-30" if dd <= 30 else "31-60" if dd <= 60 else "61-90" if dd <= 90 else "90+"
    invs.append(dict(n=n, job=job, d=d, due=due, amt=amt, paid=paid, bal=bal, status=status, bucket=bucket))
exp = {}
def E(ref, val): exp[ref] = val
J = C["Jobs"]; jt = M["jobs_total_row"]; tot = dict.fromkeys("KLMPQRST", 0.0); jobrow = {}
for i, (jid, client, desc, status, start, tend, pct, contract) in enumerate(SD.JOBS):
    r = R1 + i; jobrow[jid] = r
    H_ = esum(jid, "price"); Jc = co_price(jid); K = (H_ if contract is None else contract) + Jc
    L = esum(jid, "cost") + co_cost(jid); Mx = actual(jid)
    P = sum(x["hrs"] for x in ts if x["job"] == jid) + sum(x["qty"] for x in jc if x["job"] == jid and x["cat"] == "Labor")
    Q = sum(v["amt"] for v in invs if v["job"] == jid); Rr = sum(p["amt"] for p in pays if p["job"] == jid); T = K - max(L, Mx)
    vals = {"Estimate Price": H_, "Approved Change Orders": Jc, "Revised Contract": K, "Estimated Cost": L, "Actual Cost to Date": Mx,
            "Cost Variance": L - Mx, "% of Budget Used": (Mx / L if L else 0), "Labor Hours": P, "Invoiced": Q, "Collected": Rr, "Balance Due": Q - Rr,
            "Projected Gross Profit": T, "Projected Margin %": (T / K if K else 0), "Estimated Margin %": ((K - L) / K if K else 0), "Check": ""}
    if tend and status not in ("Complete", "Closed", "Lost"): vals["Days to Target"] = (tend - TODAY).days
    else: vals["Days to Target"] = ""
    for h, v in vals.items(): E(f"Jobs!{J[h]}{r}", v)
    for k, v in zip("KLMPQRST", [K, L, Mx, P, Q, Rr, Q - Rr, T]): tot[k] += v
for k, v in tot.items(): E(f"Jobs!{k}{jt}", v)
for i, c in enumerate(SD.CLIENTS):
    r = R1 + i; js = [j for j in SD.JOBS if j[1] == c[0]]
    E(f"Clients!{C['Clients']['Jobs']}{r}", len(js)); E(f"Clients!{C['Clients']['Balance due']}{r}", sum(exp[f"Jobs!{J['Balance Due']}{jobrow[j[0]]}"] for j in js))
for i, v in enumerate(SD.VENDORS):
    r = VEND_R1 + i; name, typ, *_, expiry, w9, elig = v
    paid = sum(x["amt"] for x in jc if x["vendor"] == name and x["d"].year == YEAR)
    E(f"Vendors & Crew!{VC['Paid YTD']}{r}", paid); E(f"Vendors & Crew!{VC['1099 Flag']}{r}", "Review for 1099" if (elig == "Yes" and paid >= 600) else "")
    E(f"Vendors & Crew!{VC['Insurance Status']}{r}", "No expiry on file" if expiry is None else ("EXPIRED" if expiry < TODAY else ("Expiring soon" if (expiry - TODAY).days <= 30 else "OK")))
for i, c in enumerate(SD.CREW):
    r = CREW_R1 + i
    E(f"Vendors & Crew!{CC['Hours YTD']}{r}", sum(x["hrs"] for x in ts if x["w"] == c[0] and x["d"].year == YEAR)); E(f"Vendors & Crew!{CC['Labor Cost YTD']}{r}", sum(x["cost"] for x in ts if x["w"] == c[0] and x["d"].year == YEAR))
Ec = C["Estimate"]
for i, e in enumerate(est):
    r = R1 + i; E(f"Estimate!{Ec['Unit Cost Used']}{r}", e["cost"] / e["qty"]); E(f"Estimate!{Ec['Price to Client']}{r}", e["price"]); E(f"Estimate!{Ec['Category']}{r}", e["cat"]); E(f"Estimate!{Ec['Check']}{r}", "")
pj = "J-1001"; sub = esum(pj, "price"); mat = esum(pj, "price", cat="Material"); total = sub + sub * cont + (sub if tax_scope == "All lines" else mat) * tax
E("Proposal!E20", sub); E("Proposal!E23", total); E("Proposal!E26", total * dep); E("Proposal!E27", total * prog); E("Proposal!E28", total - total * dep - total * prog)
for i, (cat, _) in enumerate(CATS, start=14): E(f"Proposal!E{i}", esum(pj, "price", cat=cat)); E(f"Proposal!I{i}", esum(pj, "cost", cat=cat))
lines = [e for e in est if e["job"] == pj]
for k, e in enumerate(lines, start=1): E(f"Proposal!B{30+k}", e["desc"]); E(f"Proposal!D{30+k}", e["qty"]); E(f"Proposal!E{30+k}", e["unit"]); E(f"Proposal!F{30+k}", e["price"])
E(f"Proposal!F{31+N_PROP}", sub); E(f"Proposal!D{31+N_PROP}", len(lines)); E("Proposal!I20", esum(pj, "cost")); E("Proposal!E5", "Alvarez residence")
Sc = C["Schedule"]
def workday(start, n):
    d = start; k = 1
    while k < n:
        d += dt.timedelta(days=1)
        if d.isoweekday() <= 5: k += 1
    return d
for i, (job, task, who, start, dur, pct) in enumerate(SD.SCHEDULE):
    r = R1 + i; end = workday(start, max(dur, 1)); E(f"Schedule!{Sc['End']}{r}", end)
    E(f"Schedule!{Sc['Status']}{r}", "Done" if pct >= 1 else ("Late" if TODAY > end else ("In progress" if TODAY >= start else "Not started"))); E(f"Schedule!{Sc['Check']}{r}", "")
Tc = C["Timesheet"]
for i, x in enumerate(ts):
    r = R1 + i
    E(f"Timesheet!{Tc['Rate Used']}{r}", x["rate"]); E(f"Timesheet!{Tc['Labor Cost (incl. burden)']}{r}", x["cost"]); E(f"Timesheet!{Tc['Billable Value']}{r}", x["billable"])
    E(f"Timesheet!{Tc['Week Ending']}{r}", x["we"]); E(f"Timesheet!{Tc['Category']}{r}", x["cat"]); E(f"Timesheet!{Tc['Check']}{r}", "")
week = SD.SELECTORS["Timesheet!Q3"]
for i, c in enumerate(SD.CREW):
    r = 5 + i; E(f"Timesheet!Q{r}", sum(x["hrs"] for x in ts if x["w"] == c[0] and x["we"] == week)); E(f"Timesheet!R{r}", sum(x["cost"] for x in ts if x["w"] == c[0] and x["we"] == week))
E(f"Timesheet!Q{5+N_CREW}", sum(x["hrs"] for x in ts if x["we"] == week))
Jc = C["Job Costs"]
for i, x in enumerate(jc):
    r = R1 + i; E(f"Job Costs!{Jc['Amount']}{r}", x["amt"]); E(f"Job Costs!{Jc['Category']}{r}", x["cat"]); E(f"Job Costs!{Jc['Check']}{r}", "")
Mc = C["Mileage"]
for i, x in enumerate(mil): r = R1 + i; E(f"Mileage!{Mc['Amount']}{r}", x["amt"]); E(f"Mileage!{Mc['Check']}{r}", "")
Dc = C["Daily Log"]
for i, (d, job, *_) in enumerate(SD.DAILY_LOG): E(f"Daily Log!{Dc['Hours logged (Timesheet)']}{R1+i}", sum(x["hrs"] for x in ts if x["job"] == job and x["d"] == d))
Cc = C["Change Orders"]
for i, c in enumerate(cos):
    r = R1 + i; E(f"Change Orders!{Cc['Price to Client']}{r}", c["price"]); E(f"Change Orders!{Cc['Added Profit']}{r}", c["price"] - c["cost"]); E(f"Change Orders!{Cc['Check']}{r}", "")
co2 = cos[1]; orig = 42500
E("CO Form!E14", co2["price"]); E("CO Form!E15", orig); E("CO Form!E16", co_price("J-1001") - co2["price"]); E("CO Form!E17", orig + co_price("J-1001")); E("CO Form!E18", "Approved")
E("CO Form!E13", SD.JOBS[0][5] + dt.timedelta(days=co2["sched"]))
Ic = C["Invoices"]
for i, v in enumerate(invs):
    r = R1 + i
    E(f"Invoices!{Ic['Paid to Date']}{r}", v["paid"]); E(f"Invoices!{Ic['Balance']}{r}", v["bal"]); E(f"Invoices!{Ic['Status']}{r}", v["status"]); E(f"Invoices!{Ic['Aging Bucket']}{r}", v["bucket"]); E(f"Invoices!{Ic['Check']}{r}", "")
v = [x for x in invs if x["n"] == "INV-2026-023"][0]
E("Invoice Print!E16", v["amt"]); E("Invoice Print!E17", -v["paid"]); E("Invoice Print!E18", v["bal"]); E("Invoice Print!E34", v["bal"]); E("Invoice Print!E33", "INV-2026-023")
E("Invoice Print!E21", exp[f"Jobs!{J['Revised Contract']}{jobrow['J-1001']}"]); E("Invoice Print!E24", exp[f"Jobs!{J['Revised Contract']}{jobrow['J-1001']}"] - exp[f"Jobs!{J['Invoiced']}{jobrow['J-1001']}"])
for i in range(len(SD.PAYMENTS)): E(f"Payments!{C['Payments']['Check']}{R1+i}", "")
jp = M["jpr"]
for i, (cat, _) in enumerate(CATS):
    r = jp["cat_r1"] + i; E(f"Job Profit Report!D{r}", esum(pj, "cost", cat=cat)); E(f"Job Profit Report!E{r}", actual(pj, cat=cat)); E(f"Job Profit Report!H{r}", esum(pj, "price", cat=cat))
E(f"Job Profit Report!E{jp['cat_total']}", actual(pj))
E("Job Profit Report!F9", sum(e["qty"] for e in est if e["job"] == pj and e["cat"] == "Labor" and e["unit"] == "hr")); E("Job Profit Report!G9", exp[f"Jobs!{J['Labor Hours']}{jobrow[pj]}"])
E("Job Profit Report!G8", (exp[f"Jobs!{J['Projected Gross Profit']}{jobrow[pj]}"] - exp[f"Jobs!{J['Revised Contract']}{jobrow[pj]}"] * overhead) / exp[f"Jobs!{J['Revised Contract']}{jobrow[pj]}"])
for i, code in enumerate(codes):
    r = jp["code_r1"] + i; E(f"Job Profit Report!D{r}", esum(pj, "cost", code=code)); E(f"Job Profit Report!E{r}", actual(pj, code=code))
E(f"Job Profit Report!D{jp['total_row']}", esum(pj, "cost") + co_cost(pj)); E(f"Job Profit Report!E{jp['total_row']}", actual(pj))
E(f"Job Profit Report!D{jp['total_row']+2}", sum(x["amt"] for x in jc if x["job"] == pj and x["paid"] == "No")); E(f"Job Profit Report!D{jp['total_row']+3}", sum(c["price"] for c in cos if c["job"] == pj and c["status"] == "Pending"))
rp = M["rep"]
for i, b in enumerate(["Current", "1-30", "31-60", "61-90", "90+"]):
    r = rp["aging_r1"] + i; E(f"Reports!C{r}", sum(x["bal"] for x in invs if x["bucket"] == b)); E(f"Reports!D{r}", sum(1 for x in invs if x["bucket"] == b))
E(f"Reports!C{rp['forecast_r1']}", sum(x["bal"] for x in invs if x["due"] < TODAY and x["bal"] > 0))
for m in range(6):
    r = rp["forecast_r1"] + 1 + m; y, mo = TODAY.year, TODAY.month + m
    while mo > 12: mo -= 12; y += 1
    m_start = dt.date(y, mo, 1); nxt = dt.date(y + (mo == 12), (mo % 12) + 1, 1); m_end = nxt - dt.timedelta(days=1); lo = max(m_start, TODAY)
    E(f"Reports!C{r}", sum(x["bal"] for x in invs if lo <= x["due"] <= m_end and x["bal"] > 0))
for i, s in enumerate(STATUS_LIST):
    r = rp["pipeline_r1"] + i; js = [j for j in SD.JOBS if j[3] == s]
    E(f"Reports!C{r}", len(js)); E(f"Reports!D{r}", sum(exp[f"Jobs!{J['Revised Contract']}{jobrow[j[0]]}"] for j in js))
won = sum(1 for j in SD.JOBS if j[3] in ("Approved", "In Progress", "Complete", "Closed")); lost = sum(1 for j in SD.JOBS if j[3] == "Lost")
E(f"Reports!{rp['winrate']}", won / (won + lost))
for i, code in enumerate(codes):
    r = rp["cc_r1"] + i; E(f"Reports!C{r}", sum(e["cost"] for e in est if e["code"] == code))
    E(f"Reports!D{r}", sum(x["amt"] for x in jc if x["code"] == code) + sum(x["cost"] for x in ts if x["code"] == code) + sum(x["amt"] for x in mil if x["code"] == code))
E("Tax Summary!C7", sum(p["amt"] for p in pays if p["d"].year == YEAR)); E("Tax Summary!C8", sum(v["amt"] for v in invs if v["d"].year == YEAR))
for i, (cat, _) in enumerate(CATS, start=12):
    E(f"Tax Summary!C{i}", sum(x["amt"] for x in jc if x["cat"] == cat and x["d"].year == YEAR)); E(f"Tax Summary!D{i}", sum(x["cost"] for x in ts if x["cat"] == cat and x["d"].year == YEAR)); E(f"Tax Summary!E{i}", sum(x["amt"] for x in mil if x["cat"] == cat and x["d"].year == YEAR))
E("Tax Summary!C22", sum(x["miles"] for x in mil if x["d"].year == YEAR)); E("Tax Summary!C26", sum(1 for v in SD.VENDORS if v[8] == "Yes" and sum(x["amt"] for x in jc if x["vendor"] == v[0] and x["d"].year == YEAR) >= 600))
E("Dashboard!B5", tot["K"]); E("Dashboard!E5", tot["L"]); E("Dashboard!H5", tot["M"]); E("Dashboard!K5", tot["T"])
E("Dashboard!B8", tot["T"] / tot["K"]); E("Dashboard!E8", (tot["T"] - tot["K"] * overhead) / tot["K"]); E("Dashboard!H8", tot["Q"]); E("Dashboard!K8", tot["R"])
E("Dashboard!B11", tot["S"]); E("Dashboard!E11", sum(x["bal"] for x in invs if x["status"] == "Overdue"))
act = [j for j in SD.JOBS if j[3] in ("Approved", "In Progress")]
E("Dashboard!H11", sum(exp[f"Jobs!{J['Revised Contract']}{jobrow[j[0]]}"] - exp[f"Jobs!{J['Invoiced']}{jobrow[j[0]]}"] for j in act))
E("Dashboard!K11", sum(x["amt"] for x in jc if x["paid"] == "No")); E("Dashboard!B14", len(act)); E("Dashboard!E14", sum(1 for j in SD.JOBS if exp[f"Jobs!{J['% of Budget Used']}{jobrow[j[0]]}"] > 1))
E("Dashboard!H14", sum(x["hrs"] for x in ts if x["d"].year == TODAY.year and x["d"].month == TODAY.month)); E("Dashboard!K14", won / (won + lost))
mr = M["dash"]["month_hr"]
for m in range(1, 13):
    r = mr + m
    E(f"Dashboard!C{r}", sum(v["amt"] for v in invs if v["d"].year == YEAR and v["d"].month == m))
    E(f"Dashboard!D{r}", sum(p["amt"] for p in pays if p["d"].year == YEAR and p["d"].month == m))
    E(f"Dashboard!E{r}", sum(x["amt"] for x in jc if x["d"].year == YEAR and x["d"].month == m) + sum(x["cost"] for x in ts if x["d"].year == YEAR and x["d"].month == m) + sum(x["amt"] for x in mil if x["d"].year == YEAR and x["d"].month == m))
    E(f"Dashboard!H{r}", sum(x["hrs"] for x in ts if x["d"].year == YEAR and x["d"].month == m))
hc = M["hc"]
E(f"Health Check!{hc['total_issues']}", 0); E(f"Health Check!{hc['diag_total']}", 0)
E(f"Health Check!C{hc['info_first']+1}", sum(1 for v in invs if v["status"] == "Overdue")); E(f"Health Check!C{hc['info_first']+2}", sum(1 for k in exp if k.startswith("Schedule!") and exp[k] == "Late"))
E(f"Health Check!C{hc['info_first']+3}", sum(1 for v in SD.VENDORS if v[6] is not None and v[6] < TODAY))
refs = list(exp.keys())
out = subprocess.run(["osascript", os.path.join(HERE, "excel_recalc.applescript"), XLSX] + refs, capture_output=True, text=True)
if out.returncode != 0: print("AppleScript failed:", out.stderr); sys.exit(2)
got = {}
for line in out.stdout.strip().split("\n"):
    if "\t" in line: k, v = line.split("\t", 1); got[k] = v
def parse_date(s):
    for f in ("%A, %B %d, %Y at %I:%M:%S %p", "%Y-%m-%d"):
        try: return dt.datetime.strptime(s, f).date()
        except Exception: pass
    return None
fail = 0; shown = 0
for k in refs:
    e = exp[k]; g = got.get(k)
    if isinstance(e, dt.date): ok = parse_date(g or "") == e
    elif isinstance(e, str): ok = (g == e)
    else:
        try: ok = abs(float(g) - float(e)) < 0.005
        except Exception: ok = False
    if not ok:
        fail += 1
        if shown < 40: print("FAIL", k, "expected", e, "excel", repr(g)); shown += 1
print(f"\n{len(refs) - fail}/{len(refs)} checks passed")
sys.exit(1 if fail else 0)
