"""Settings, Clients, Vendors & Crew, Cost Library (v3)."""
from openpyxl.formatting.rule import CellIsRule
from layout import *

def build_settings(wb, blank):
    st = wb.create_sheet("Settings")
    title(st, "Settings", "Light-blue cells are inputs. Everything here feeds dropdowns, markups, documents and reports on the other sheets.", TAB["setup"])
    for col_, w in zip("ABCDEFGHIJKLM", [54, 44, 58, 3, 10, 34, 18, 4, 14, 12, 17, 18, 4]):
        st.column_dimensions[col_].width = w
    header(st, 4, ["Company & defaults", "Value"])
    company = [("Company name", "Your Company LLC" if blank else "Northline Remodeling LLC", None),
               ("Address", "" if blank else "410 Harbor Rd, Dover, DE 19901", None),
               ("Phone / email", "" if blank else "(555) 010-2200 · office@northline.example", None),
               ("License #", "" if blank else "DE-HIC-2026-0417", None),
               ("Default contingency %", 0.05, PCT), ("Sales tax %", 0.0, PCT),
               ("Deposit % of proposal total", 0.30, PCT), ("Progress payment % (mid-job)", 0.40, PCT),
               ("Overdue grace (days past due)", 0, None), ("Reporting / tax year", 2026, None),
               ("Mileage rate ($ per mile)", 0.70, MONEY), ("Proposal valid for (days)", 30, None),
               ("Overhead % of revenue (for net margin)", 0.10, PCT),
               ("Labor burden % (payroll taxes, insurance, added to crew cost)", 0.0 if blank else 0.12, PCT),
               ("Sales tax applies to", "Material only", None)]
    for i, (k, v, f) in enumerate(company, start=5):
        st.cell(row=i, column=1, value=k).border = BORDER
        inp(st, i, 2, v, f)
    dv_list(st, f'"{",".join(TAX_SCOPE_LIST)}"', "B19")
    notes = {15: "Check the current IRS standard mileage rate each January.", 18: "Applied on the Timesheet: Labor cost = hours × rate × (1 + burden).",
             19: "Material only = tax on Material lines. All lines = tax on the whole subtotal.", 17: "Used for 'margin after overhead' on the Dashboard and Job Profit Report."}
    for r, t in notes.items(): st.cell(row=r, column=3, value=t).font = NOTE
    header(st, DOC_HDR, ["Document text", "Value"], height=18)
    texts = [("Proposal terms", "This proposal is valid until the date shown. Work begins on receipt of the deposit. Changes to scope are priced as written change orders before work proceeds. Balance is due on completion."),
             ("Invoice payment instructions", "Pay by check to the company name above, or by bank transfer. Please reference the invoice number."),
             ("Invoice / statement footer", "Thank you for your business."),
             ("Change order terms", "The client authorizes the added work and price above. The revised contract total and any schedule change apply from the date signed."),
             ("Proposal exclusions & assumptions", "Excludes permits not listed, hazardous-material remediation, structural repairs discovered after demolition, and work outside the described scope. Pricing assumes access to water and power on site.")]
    for i, (k, v) in enumerate(texts, start=DOC_R1):
        st.cell(row=i, column=1, value=k).border = BORDER
        x = inp(st, i, 2, v); x.alignment = Alignment(wrap_text=True, vertical="top"); st.row_dimensions[i].height = max(34, 16 * (len(v) // 52 + 1) + 6)
    header(st, MARKUP_R1 - 1, ["Cost category", "Default markup %"], height=18)
    for i, (k, v) in enumerate(CATS, start=MARKUP_R1):
        st.cell(row=i, column=1, value=k).border = BORDER
        inp(st, i, 2, v, PCT)
    st.cell(row=MARKUP_R2 + 1, column=1, value="Markup is applied to cost: Price = Cost × (1 + markup). Override per line on the Estimate.").font = NOTE
    st.cell(row=DOC_R1 + 5, column=1, value="Logo: on Proposal, Invoice Print, CO Form or Client Statement use Insert → Pictures and place it top-right.").font = NOTE
    header(st, 4, ["Code", "Cost code name", "Category"], start_col=5)
    codes = [("01", "General conditions / supervision", "Labor"), ("02", "Demolition & disposal", "Labor"),
             ("03", "Site work & excavation", "Subcontractor"), ("04", "Concrete & foundation", "Subcontractor"),
             ("05", "Framing labor", "Labor"), ("06", "Lumber & framing materials", "Material"),
             ("07", "Roofing", "Subcontractor"), ("08", "Windows & doors", "Material"),
             ("09", "Exterior finishes / siding", "Subcontractor"), ("10", "Plumbing", "Subcontractor"),
             ("11", "Electrical", "Subcontractor"), ("12", "HVAC", "Subcontractor"),
             ("13", "Insulation & drywall", "Subcontractor"), ("14", "Interior finish labor", "Labor"),
             ("15", "Paint & wall finishes", "Material"), ("16", "Flooring & tile", "Material"),
             ("17", "Cabinets & countertops", "Material"), ("18", "Fixtures & appliances", "Material"),
             ("19", "Equipment rental", "Equipment"), ("20", "Permits & inspections", "Permit/Fee"),
             ("21", "Dumpster & cleanup", "Other"), ("22", "Fuel & mileage", "Other"),
             ("23", "Warranty / callbacks", "Labor"), ("24", "Miscellaneous", "Other")]
    for i, row in enumerate(codes, start=CODE_R1):
        for j, v in enumerate(row, start=5): inp(st, i, j, v)
    dv_list(st, f"={CAT_RANGE}", f"G{CODE_R1}:G{CODE_R2}", "Category drives the default markup and the reports.")
    lists = {"Job status": STATUS_LIST, "Units": UNIT_LIST, "Payment methods": PAYM_LIST, "Vendor types": VTYPE_LIST}
    for j, (k, vals) in enumerate(lists.items(), start=9):
        c = st.cell(row=4, column=j, value=k); c.fill = HEAD_FILL; c.font = WHITE_BOLD; c.border = BORDER
        for i, v in enumerate(vals, start=5): inp(st, i, j, v)
    print_setup(st)
    return st

def build_clients(wb):
    ws = wb.create_sheet("Clients")
    title(ws, "Clients", "One row per client. Jobs, contract value, collected and balance due are calculated from the Jobs sheet. Print a statement from the Client Statement sheet.", TAB["input"])
    header(ws, 4, H["Clients"], [24, 20, 14, 26, 34, 24, 8, 18, 14, 14, 24])
    ws.freeze_panes = "B5"; c = C["Clients"]
    for r in range(R1, last("Clients") + 1):
        for h in ["Client", "Contact person", "Phone", "Email", "Billing address", "Notes"]: inp(ws, r, ord(c[h]) - 64)
        k = f"{c['Client']}{r}"
        fx(ws, r, ord(c["Jobs"]) - 64, f'=IF({k}="","",COUNTIF({rr("Jobs","Client")},{k}))')
        fx(ws, r, ord(c["Contract value (all jobs)"]) - 64, f'=IF({k}="","",SUMIFS({rr("Jobs","Revised Contract")},{rr("Jobs","Client")},{k}))', MONEY)
        fx(ws, r, ord(c["Collected"]) - 64, f'=IF({k}="","",SUMIFS({rr("Jobs","Collected")},{rr("Jobs","Client")},{k}))', MONEY)
        fx(ws, r, ord(c["Balance due"]) - 64, f'=IF({k}="","",SUMIFS({rr("Jobs","Balance Due")},{rr("Jobs","Client")},{k}))', MONEY)
        fx(ws, r, ord(c["Check"]) - 64, f'=IF({k}="","",IF(COUNTIF({CLIENT_NAMES},{k})>1,"Duplicate client",""))')
    from openpyxl.formatting.rule import FormulaRule
    ws.conditional_formatting.add(f"{c['Check']}{R1}:{c['Check']}{last('Clients')}", FormulaRule(formula=[f'{c["Check"]}{R1}<>""'], fill=RED_FILL))
    print_setup(ws, title_rows="4:4")
    return ws

def build_vendors(wb):
    ws = wb.create_sheet("Vendors & Crew")
    title(ws, "Vendors & Crew", "Crew (your workers, with cost and bill rates) feeds the Timesheet. Vendors & subs feed Job Costs, insurance alerts and the 1099 review.", TAB["input"])
    ws.cell(row=3, column=1, value="CREW").font = H2
    header(ws, 4, CREW_H, [24, 18, 14, 14, 14, 12, 16, 22, 26])
    y = S["year"]
    for r in range(CREW_R1, CREW_R2 + 1):
        for h in ["Worker", "Role", "Cost Rate ($/hr)", "Bill Rate ($/hr)", "Phone", "Notes"]:
            inp(ws, r, ord(CC[h]) - 64, fmt=MONEY if "Rate" in h else None)
        k = f"{CC['Worker']}{r}"
        d1 = f"DATE({y},1,1)"; d2 = f"DATE({y},12,31)"
        fx(ws, r, ord(CC["Hours YTD"]) - 64, f'=IF({k}="","",SUMIFS({rr("Timesheet","Hours")},{rr("Timesheet","Worker")},{k},{rr("Timesheet","Date")},">="&{d1},{rr("Timesheet","Date")},"<="&{d2}))', HOURS)
        fx(ws, r, ord(CC["Labor Cost YTD"]) - 64, f'=IF({k}="","",SUMIFS({rr("Timesheet","Labor Cost (incl. burden)")},{rr("Timesheet","Worker")},{k},{rr("Timesheet","Date")},">="&{d1},{rr("Timesheet","Date")},"<="&{d2}))', MONEY)
        fx(ws, r, ord(CC["Check"]) - 64, f'=IF({k}="","",IF(COUNTIF({CREW_NAMES},{k})>1,"Duplicate worker",IF({CC["Cost Rate ($/hr)"]}{r}="","Missing cost rate","")))')
    ws.cell(row=VEND_HDR - 1, column=1, value="VENDORS & SUBCONTRACTORS").font = H2
    header(ws, VEND_HDR, VEND_H, [24, 16, 20, 14, 28, 12, 16, 18, 12, 12, 14, 14, 18, 22, 26])
    for r in range(VEND_R1, VEND_R2 + 1):
        for h in ["Vendor / Sub", "Type", "Trade", "Phone", "Email", "License #", "Insurance Expiry", "W-9 on file", "1099-eligible", "Notes"]:
            inp(ws, r, ord(VC[h]) - 64, fmt=DATEF if h == "Insurance Expiry" else None)
        k = f"{VC['Vendor / Sub']}{r}"; e = f"{VC['Insurance Expiry']}{r}"
        fx(ws, r, ord(VC["Insurance Status"]) - 64, f'=IF({k}="","",IF({e}="","No expiry on file",IF({e}<TODAY(),"EXPIRED",IF({e}-TODAY()<=30,"Expiring soon","OK"))))')
        d1 = f"DATE({y},1,1)"; d2 = f"DATE({y},12,31)"
        fx(ws, r, ord(VC["Paid YTD"]) - 64, f'=IF({k}="","",SUMIFS({rr("Job Costs","Amount")},{rr("Job Costs","Vendor")},{k},{rr("Job Costs","Date")},">="&{d1},{rr("Job Costs","Date")},"<="&{d2}))', MONEY)
        fx(ws, r, ord(VC["Unpaid Balance"]) - 64, f'=IF({k}="","",SUMIFS({rr("Job Costs","Amount")},{rr("Job Costs","Vendor")},{k},{rr("Job Costs","Paid?")},"No"))', MONEY)
        fx(ws, r, ord(VC["1099 Flag"]) - 64, f'=IF({k}="","",IF(AND({VC["1099-eligible"]}{r}="Yes",{VC["Paid YTD"]}{r}>=600),"Review for 1099",""))')
        fx(ws, r, ord(VC["Check"]) - 64, f'=IF({k}="","",IF(COUNTIF({VEND_NAMES},{k})>1,"Duplicate vendor",""))')
    dv_list(ws, f"={VTYPE_RANGE}", f"{VC['Type']}{VEND_R1}:{VC['Type']}{VEND_R2}")
    dv_list(ws, '"Yes,No"', f"{VC['W-9 on file']}{VEND_R1}:{VC['W-9 on file']}{VEND_R2}")
    dv_list(ws, '"Yes,No"', f"{VC['1099-eligible']}{VEND_R1}:{VC['1099-eligible']}{VEND_R2}")
    s = f"{VC['Insurance Status']}{VEND_R1}:{VC['Insurance Status']}{VEND_R2}"
    ws.conditional_formatting.add(s, CellIsRule(operator="equal", formula=['"EXPIRED"'], fill=RED_FILL))
    ws.conditional_formatting.add(s, CellIsRule(operator="equal", formula=['"Expiring soon"'], fill=YEL_FILL))
    ws.conditional_formatting.add(s, CellIsRule(operator="equal", formula=['"OK"'], fill=GRN_FILL))
    ws.conditional_formatting.add(f"{VC['1099 Flag']}{VEND_R1}:{VC['1099 Flag']}{VEND_R2}", CellIsRule(operator="equal", formula=['"Review for 1099"'], fill=YEL_FILL))
    from openpyxl.formatting.rule import FormulaRule
    ws.conditional_formatting.add(f"{VC['Unpaid Balance']}{VEND_R1}:{VC['Unpaid Balance']}{VEND_R2}", FormulaRule(formula=[f'AND({VC["Unpaid Balance"]}{VEND_R1}<>"",{VC["Unpaid Balance"]}{VEND_R1}>0)'], fill=YEL_FILL))
    ws.cell(row=VEND_R2 + 2, column=1, value="1099 flag is informational only (US vendors paid $600+ in the year and marked eligible). Confirm requirements with your tax professional.").font = NOTE
    print_setup(ws)
    return ws

def build_library(wb):
    ws = wb.create_sheet("Cost Library")
    title(ws, "Cost Library (price book)", "Save the items you estimate repeatedly. On the Estimate sheet, pick a Library Item and the unit and unit cost fill in unless you type your own.", TAB["input"])
    header(ws, 4, H["Cost Library"], [40, 10, 8, 14, 14, 12, 30]); ws.freeze_panes = "B5"; c = C["Cost Library"]
    for r in range(R1, last("Cost Library") + 1):
        for h in ["Item", "Cost Code", "Unit", "Unit Cost", "Notes"]: inp(ws, r, ord(c[h]) - 64, fmt=MONEY if h == "Unit Cost" else None)
        fx(ws, r, ord(c["Category"]) - 64, f'=IF({c["Cost Code"]}{r}="","",IFERROR(VLOOKUP({c["Cost Code"]}{r},{CODE_TABLE},3,FALSE),"Other"))')
        fx(ws, r, ord(c["Used in estimates"]) - 64, f'=IF({c["Item"]}{r}="","",COUNTIF({rr("Estimate","Library Item")},{c["Item"]}{r}))')
    dv_list(ws, f"={CODE_RANGE}", f"{c['Cost Code']}{R1}:{c['Cost Code']}{last('Cost Library')}", "Cost codes are defined on Settings.")
    dv_list(ws, f"={UNIT_RANGE}", f"{c['Unit']}{R1}:{c['Unit']}{last('Cost Library')}")
    print_setup(ws, title_rows="4:4")
    return ws
