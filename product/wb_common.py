"""Shared styles, helpers and layout constants for the workbook generator (v2)."""
import os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

NAVY = "1F3A5F"; TEAL = "2A7F8E"
INPUT_FILL = PatternFill("solid", fgColor="EAF3FF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="D9E2EF")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
KPI_FILL = PatternFill("solid", fgColor="EEF6F7")
def _cf(c): return PatternFill(start_color=c, end_color=c, fill_type="solid")  # both colours: required for conditional-format fills
RED_FILL = _cf("F8D7DA"); YEL_FILL = _cf("FFF3CD"); GRN_FILL = _cf("D4EDDA"); GRAY_FILL = _cf("EDEDED")
BAR_FILL = _cf("5B9BD5"); BAR_DONE_FILL = _cf("1F3A5F"); TODAY_FILL = _cf("FFE699")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=16, color=NAVY)
H2 = Font(bold=True, size=12, color=NAVY)
NOTE = Font(italic=True, color="666666", size=9)
BLUE_INPUT = Font(color="1F4E9A")
LINK = Font(color="1F4E9A", underline="single")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '"$"#,##0.00;[Red]-"$"#,##0.00'
MONEY0 = '"$"#,##0;[Red]-"$"#,##0'
PCT = '0.0%'
DATEF = 'yyyy-mm-dd'
HOURS = '#,##0.0'

# Capacities (first data row is always 5)
N_JOBS, N_EST, N_TS, N_COST, N_MIL, N_CO, N_INV, N_PAY = 40, 400, 1000, 1000, 300, 100, 200, 300
N_SCH, N_CLI, N_CREW, N_VEND, N_LIB, N_LOG, N_CODES, N_PROP = 200, 100, 30, 100, 200, 300, 24, 40
GANTT_DAYS = 84
if os.environ.get("MINI"):  # small-capacity variant used only for the Google Sheets round-trip test
    N_JOBS, N_EST, N_TS, N_COST, N_MIL, N_CO, N_INV, N_PAY = 8, 40, 30, 25, 8, 8, 8, 8
    N_SCH, N_CLI, N_CREW, N_VEND, N_LIB, N_LOG, N_PROP = 20, 8, 6, 14, 14, 8, 16
    GANTT_DAYS = 21
TAB = {"setup": "7F7F7F", "input": "4A86C8", "report": "3A9D5D"}

def q(sheet):
    return f"'{sheet}'" if " " in sheet or "&" in sheet else sheet

def rng(sheet, col, r1, r2):
    return f"{q(sheet)}!${col}${r1}:${col}${r2}"

def cell(sheet, col, r):
    return f"{q(sheet)}!${col}${r}"

def colmap(headers):
    return {h: get_column_letter(i + 1) for i, h in enumerate(headers)}

def header(ws, row, headers, widths=None, start_col=1, fill=HEAD_FILL, font=WHITE_BOLD, height=30):
    for i, h in enumerate(headers):
        if h is None: continue
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.fill = fill; c.font = font; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(start_col + i)].width = w

def title(ws, text, sub=None, color=None):
    ws["A1"] = text; ws["A1"].font = TITLE
    if sub:
        ws["A2"] = sub; ws["A2"].font = NOTE
    ws.sheet_view.showGridLines = False
    if color: ws.sheet_properties.tabColor = color

def style_range(ws, r1, r2, c1, c2, fill=None, font=None, fmt=None, border=True, align=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = ws.cell(row=r, column=c)
            if fill: x.fill = fill
            if font: x.font = font
            if fmt: x.number_format = fmt
            if border: x.border = BORDER
            if align: x.alignment = align

def inp(ws, r, c, value=None, fmt=None):
    x = ws.cell(row=r, column=c)
    if value is not None: x.value = value
    x.fill = INPUT_FILL; x.font = BLUE_INPUT; x.border = BORDER; x.protection = Protection(locked=False)
    if fmt: x.number_format = fmt
    return x

def fx(ws, r, c, formula, fmt=None, bold=False):
    x = ws.cell(row=r, column=c, value=formula); x.border = BORDER
    if fmt: x.number_format = fmt
    if bold: x.font = BOLD
    return x

def dv_list(ws, formula, cells, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
    if prompt:
        dv.showInputMessage = True; dv.promptTitle = "Pick from the list"; dv.prompt = prompt
    ws.add_data_validation(dv); dv.add(cells)

def print_setup(ws, landscape=True, title_rows=None, area=None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "&A  —  page &P of &N"
    ws.oddFooter.center.size = 8
    if title_rows: ws.print_title_rows = title_rows
    if area: ws.print_area = area

def link(ws, r, c, text, sheet):
    x = ws.cell(row=r, column=c, value=text)
    x.hyperlink = f"#{q(sheet)}!A1"; x.font = LINK
    return x
