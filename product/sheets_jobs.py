"""Jobs, Estimate, Proposal."""
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from layout import *

def ci(name, h): return ord(C[name][h]) - 64

def build_jobs(wb):
    ws = wb.create_sheet("Jobs"); n = "Jobs"; c = C[n]
    title(ws, "Jobs", "One row per job. Costs, hours, invoices and payments flow in from the other sheets. Contract Value is optional: leave blank to use the Estimate Price.", TAB["input"])
    header(ws, 4, H[n], [12, 22, 34, 12, 12, 12, 11, 14, 14, 14, 14, 14, 14, 14, 11, 10, 14, 14, 14, 14, 11, 11, 10, 26, 30])
    ws.freeze_panes = "C5"
    ws["A3"] = f'="Suggested next Job ID: J-"&(1000+COUNTA(A{R1}:A{last(n)})+1)'; ws["A3"].font = NOTE
    for r in range(R1, last(n) + 1):
        for h in ["Job ID", "Client", "Job address / description", "Status", "Start date", "Target end", "% Complete (manual)", "Contract Value (signed)", "Notes"]:
            inp(ws, r, ci(n, h), fmt={"Start date": DATEF, "Target end": DATEF, "% Complete (manual)": PCT, "Contract Value (signed)": MONEY}.get(h))
        A = f"$A{r}"; g = lambda h: f"{c[h]}{r}"
        f = {
            "Estimate Price": f'=IF({A}="","",SUMIFS({rr("Estimate","Price to Client")},{rr("Estimate","Job ID")},{A}))',
            "Approved Change Orders": f'=IF({A}="","",SUMIFS({rr("Change Orders","Price to Client")},{rr("Change Orders","Job ID")},{A},{rr("Change Orders","Status")},"Approved"))',
            "Revised Contract": f'=IF({A}="","",IF({g("Contract Value (signed)")}="",{g("Estimate Price")},{g("Contract Value (signed)")})+{g("Approved Change Orders")})',
            "Estimated Cost": f'=IF({A}="","",SUMIFS({rr("Estimate","Est. Cost")},{rr("Estimate","Job ID")},{A})+SUMIFS({rr("Change Orders","Added Cost (budget)")},{rr("Change Orders","Job ID")},{A},{rr("Change Orders","Status")},"Approved"))',
            "Actual Cost to Date": f'=IF({A}="","",SUMIFS({rr("Job Costs","Amount")},{rr("Job Costs","Job ID")},{A})+SUMIFS({rr("Timesheet","Labor Cost (incl. burden)")},{rr("Timesheet","Job ID")},{A})+SUMIFS({rr("Mileage","Amount")},{rr("Mileage","Job ID")},{A}))',
            "Cost Variance": f'=IF({A}="","",{g("Estimated Cost")}-{g("Actual Cost to Date")})',
            "% of Budget Used": f'=IF({A}="","",IF({g("Estimated Cost")}=0,0,{g("Actual Cost to Date")}/{g("Estimated Cost")}))',
            "Labor Hours": f'=IF({A}="","",SUMIFS({rr("Timesheet","Hours")},{rr("Timesheet","Job ID")},{A})+SUMIFS({rr("Job Costs","Qty or Hours")},{rr("Job Costs","Job ID")},{A},{rr("Job Costs","Category")},"Labor"))',
            "Invoiced": f'=IF({A}="","",SUMIFS({rr("Invoices","Amount")},{rr("Invoices","Job ID")},{A}))',
            "Collected": f'=IF({A}="","",SUMIFS({rr("Payments","Amount")},{rr("Payments","Job ID")},{A}))',
            "Balance Due": f'=IF({A}="","",{g("Invoiced")}-{g("Collected")})',
            "Projected Gross Profit": f'=IF({A}="","",{g("Revised Contract")}-MAX({g("Estimated Cost")},{g("Actual Cost to Date")}))',
            "Projected Margin %": f'=IF({A}="","",IF({g("Revised Contract")}=0,0,{g("Projected Gross Profit")}/{g("Revised Contract")}))',
            "Estimated Margin %": f'=IF({A}="","",IF({g("Revised Contract")}=0,0,({g("Revised Contract")}-{g("Estimated Cost")})/{g("Revised Contract")}))',
            "Days to Target": f'=IF(OR({A}="",{g("Target end")}="",{g("Status")}="Complete",{g("Status")}="Closed",{g("Status")}="Lost"),"",{g("Target end")}-TODAY())',
            "Check": f'=IF({A}="","",IF(COUNTIF({JOB_IDS},{A})>1,"Duplicate Job ID",IF(AND({g("Client")}<>"",COUNTIF({CLIENT_NAMES},{g("Client")})=0),"Client not in Clients sheet",IF({g("Status")}="","Missing status",IF(AND(OR({g("Status")}="Approved",{g("Status")}="In Progress"),{g("Estimate Price")}=0,{g("Contract Value (signed)")}=""),"Active job with no estimate or contract value","")))))',
        }
        for h, formula in f.items():
            fmt = PCT if "%" in h else (None if h in ("Days to Target", "Check") else (HOURS if h == "Labor Hours" else MONEY))
            fx(ws, r, ci(n, h), formula, fmt)
    tr = jobs_total_row()
    ws.cell(row=tr, column=1, value="TOTAL")
    for h in ["Estimate Price", "Contract Value (signed)", "Approved Change Orders", "Revised Contract", "Estimated Cost", "Actual Cost to Date", "Cost Variance", "Labor Hours", "Invoiced", "Collected", "Balance Due", "Projected Gross Profit"]:
        L = c[h]; fx(ws, tr, ci(n, h), f"=SUM({L}{R1}:{L}{last(n)})", HOURS if h == "Labor Hours" else MONEY, bold=True)
    fx(ws, tr, ci(n, "Projected Margin %"), f'=IF({c["Revised Contract"]}{tr}=0,0,{c["Projected Gross Profit"]}{tr}/{c["Revised Contract"]}{tr})', PCT, True)
    fx(ws, tr, ci(n, "% of Budget Used"), f'=IF({c["Estimated Cost"]}{tr}=0,0,{c["Actual Cost to Date"]}{tr}/{c["Estimated Cost"]}{tr})', PCT, True)
    fx(ws, tr, ci(n, "Estimated Margin %"), f'=IF({c["Revised Contract"]}{tr}=0,0,({c["Revised Contract"]}{tr}-{c["Estimated Cost"]}{tr})/{c["Revised Contract"]}{tr})', PCT, True)
    style_range(ws, tr, tr, 1, len(H[n]), fill=TOTAL_FILL, font=BOLD)
    dv_list(ws, f"={STATUS_RANGE}", f"{c['Status']}{R1}:{c['Status']}{last(n)}", "Lead → Quoted → Approved → In Progress → Complete → Closed. Lost closes a quote.")
    dv_list(ws, f"={CLIENT_NAMES}", f"{c['Client']}{R1}:{c['Client']}{last(n)}", "Add the client on the Clients sheet first.")
    ws.conditional_formatting.add(f"{c['% of Budget Used']}{R1}:{c['% of Budget Used']}{last(n)}", FormulaRule(formula=[f'AND({c["% of Budget Used"]}{R1}<>"",{c["% of Budget Used"]}{R1}>1)'], fill=RED_FILL))
    ws.conditional_formatting.add(f"{c['Projected Margin %']}{R1}:{c['Projected Margin %']}{last(n)}", FormulaRule(formula=[f'AND({c["Projected Margin %"]}{R1}<>"",{c["Projected Margin %"]}{R1}<0.1)'], fill=YEL_FILL))
    ws.conditional_formatting.add(f"{c['Check']}{R1}:{c['Check']}{last(n)}", FormulaRule(formula=[f'{c["Check"]}{R1}<>""'], fill=RED_FILL))
    print_setup(ws, title_rows="4:4")
    return ws

def build_estimate(wb):
    ws = wb.create_sheet("Estimate"); n = "Estimate"; c = C[n]
    title(ws, "Estimate", "One line per item. Pick a Cost Code (drives category and markup) and optionally a Library Item (fills unit and unit cost unless you type your own).", TAB["input"])
    header(ws, 4, H[n], [12, 10, 28, 36, 8, 8, 12, 14, 12, 9, 14, 10, 14, 10, 9, 22, 26])
    ws.freeze_panes = "E5"
    for r in range(R1, last(n) + 1):
        for h in ["Job ID", "Cost Code", "Library Item", "Description", "Qty", "Unit", "Unit Cost", "Markup Override", "Notes"]:
            inp(ws, r, ci(n, h), fmt={"Unit Cost": MONEY, "Markup Override": PCT}.get(h))
        A = f"$A{r}"; g = lambda h: f"{c[h]}{r}"
        fx(ws, r, ci(n, "Category"), f'=IF({g("Cost Code")}="","",IFERROR(VLOOKUP({g("Cost Code")},{CODE_TABLE},3,FALSE),"Other"))')
        fx(ws, r, ci(n, "Unit Cost Used"), f'=IF({A}="","",IF({g("Unit Cost")}<>"",{g("Unit Cost")},IFERROR(VLOOKUP({g("Library Item")},{LIB_TABLE},4,FALSE),0)))', MONEY)
        fx(ws, r, ci(n, "Unit Used"), f'=IF({A}="","",IF({g("Unit")}<>"",{g("Unit")},IFERROR(VLOOKUP({g("Library Item")},{LIB_TABLE},3,FALSE),"")))')
        fx(ws, r, ci(n, "Est. Cost"), f'=IF({A}="","",{g("Qty")}*{g("Unit Cost Used")})', MONEY)
        fx(ws, r, ci(n, "Markup Used"), f'=IF({A}="","",IF({g("Markup Override")}="",IFERROR(VLOOKUP({g("Category")},{MARKUP_TABLE},2,FALSE),0),{g("Markup Override")}))', PCT)
        fx(ws, r, ci(n, "Price to Client"), f'=IF({A}="","",{g("Est. Cost")}*(1+{g("Markup Used")}))', MONEY)
        fx(ws, r, ci(n, "Proposal Line #"), f'=IF(AND({A}<>"",{A}={PROPOSAL_SEL}),COUNTIF($A${R1}:$A{r},{PROPOSAL_SEL}),"")')
        fx(ws, r, ci(n, "Check"), f'=IF({A}="","",IF(COUNTIF({JOB_IDS},{A})=0,"Unknown Job ID",IF({g("Cost Code")}="","Missing cost code",IF(COUNTIF({CODE_RANGE},{g("Cost Code")})=0,"Unknown cost code",IF(AND({g("Unit Cost")}="",{g("Library Item")}=""),"No unit cost",IF({g("Qty")}="","Missing qty",""))))))')
    tr = last(n) + 1
    ws.cell(row=tr, column=ci(n, "Description"), value="TOTAL (all jobs)")
    for h in ["Est. Cost", "Price to Client"]:
        L = c[h]; fx(ws, tr, ci(n, h), f"=SUM({L}{R1}:{L}{last(n)})", MONEY, True)
    style_range(ws, tr, tr, 1, len(H[n]), fill=TOTAL_FILL, font=BOLD)
    dv_list(ws, f"={JOB_IDS}", f"A{R1}:A{last(n)}", "Job IDs come from the Jobs sheet.")
    dv_list(ws, f"={CODE_RANGE}", f"{c['Cost Code']}{R1}:{c['Cost Code']}{last(n)}", "Cost code sets the category and default markup.")
    dv_list(ws, f"={LIB_NAMES}", f"{c['Library Item']}{R1}:{c['Library Item']}{last(n)}", "Optional. Fills unit and unit cost from the Cost Library unless you type your own.")
    dv_list(ws, f"={UNIT_RANGE}", f"{c['Unit']}{R1}:{c['Unit']}{last(n)}")
    ws.conditional_formatting.add(f"{c['Check']}{R1}:{c['Check']}{last(n)}", FormulaRule(formula=[f'{c["Check"]}{R1}<>""'], fill=RED_FILL))
    print_setup(ws, title_rows="4:4")
    return ws

def build_proposal(wb):
    ws = wb.create_sheet("Proposal")
    ws.sheet_properties.tabColor = TAB["report"]; ws.sheet_view.showGridLines = False
    for col_, w in zip("ABCDEFGHIJ", [9, 44, 10, 10, 16, 14, 4, 22, 16, 16]): ws.column_dimensions[col_].width = w
    e = C["Estimate"]; j = C["Jobs"]; cl = C["Clients"]
    EJ = rr("Estimate", "Job ID"); ECAT = rr("Estimate", "Category"); ECOST = rr("Estimate", "Est. Cost"); EPRICE = rr("Estimate", "Price to Client")
    ws["A1"] = "PROPOSAL"; ws["A1"].font = Font(bold=True, size=22, color=NAVY)
    ws["H2"] = "Job ID"; ws["H2"].font = BOLD; inp(ws, 2, 9); dv_list(ws, f"={JOB_IDS}", "I2", "Pick the job to build the proposal for.")
    ws["H3"] = "← pick a job; the proposal at left fills in. Columns H–J never print."; ws["H3"].font = NOTE
    JOBROW = f"MATCH($I$2,{JOB_IDS},0)"
    def jobf(h): return f'IFERROR(INDEX({rr("Jobs", h)},{JOBROW}),"")'
    CLIENT = jobf("Client")
    def clif(h): return f'IFERROR(INDEX({rr("Clients", h)},MATCH({CLIENT},{CLIENT_NAMES},0)),"")'
    ws["B4"] = "FROM"; ws["E4"] = "PREPARED FOR"
    for a in ("B4", "E4"): ws[a].font = Font(bold=True, color=TEAL)
    ws["B5"] = f"={S['company']}&\"\""; ws["B5"].font = BOLD
    ws["B6"] = f"={S['address']}&\"\""; ws["B7"] = f"={S['phone']}&\"\""; ws["B8"] = f'=IF({S["license"]}="","","License # "&{S["license"]})'
    ws["E5"] = f"={CLIENT}"; ws["E5"].font = BOLD
    ws["E6"] = f'={clif("Contact person")}&IF({clif("Phone")}="","","  ·  "&{clif("Phone")})'; ws["E7"] = f"={clif('Billing address')}"
    ws["E8"] = f"={clif('Email')}"
    for r_ in range(5, 9): ws.merge_cells(start_row=r_, start_column=5, end_row=r_, end_column=6)
    ws["B10"] = "Project"; ws["B10"].font = BOLD; ws["B11"] = f"={jobf('Job address / description')}"
    ws["E9"] = "Proposal #"; ws["E9"].font = BOLD; ws["F9"] = '=IF($I$2="","","P-"&$I$2)'
    ws["E10"] = "Proposal date"; ws["E10"].font = BOLD; ws["F10"] = "=TODAY()"; ws["F10"].number_format = DATEF
    ws["E11"] = "Valid until"; ws["E11"].font = BOLD; ws["F11"] = f"=TODAY()+{S['valid_days']}"; ws["F11"].number_format = DATEF
    header(ws, 13, [None, "Scope summary", "", "", "Price", "% of total"])
    ws.cell(row=13, column=1).fill = PatternFill(None)
    ws.merge_cells("B13:D13")
    for i, (cat, _) in enumerate(CATS, start=14):
        ws.cell(row=i, column=2, value=cat).border = BORDER
        fx(ws, i, 5, f'=SUMIFS({EPRICE},{EJ},$I$2,{ECAT},B{i})', MONEY)
        fx(ws, i, 6, f'=IF($E$20=0,0,E{i}/$E$20)', PCT)
        for cc in (3, 4): ws.cell(row=i, column=cc).border = BORDER
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    for i in list(range(20, 24)) + [26, 27, 28]: ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    ws["B20"] = "Subtotal"; ws["E20"] = "=SUM(E14:E19)"
    ws["B21"] = "Contingency"; ws["F21"] = f"={S['contingency']}"; ws["E21"] = "=E20*F21"
    ws["B22"] = f'=IF({S["tax_scope"]}="All lines","Sales tax","Sales tax on materials")'; ws["F22"] = f"={S['tax']}"; ws["E22"] = f'=IF({S["tax_scope"]}="All lines",E20*F22,E15*F22)'
    ws["B23"] = "PROPOSAL TOTAL"; ws["E23"] = "=E20+E21+E22"
    for r in range(20, 24):
        ws.cell(row=r, column=5).number_format = MONEY; ws.cell(row=r, column=6).number_format = PCT
        style_range(ws, r, r, 2, 6, fill=TOTAL_FILL if r in (20, 23) else None, font=BOLD if r in (20, 23) else None)
    ws["B25"] = "Payment schedule"; ws["B25"].font = Font(bold=True, color=TEAL)
    ws["B26"] = "Deposit at signing"; ws["F26"] = f"={S['deposit']}"; ws["E26"] = "=E23*F26"
    ws["B27"] = "Progress payment (mid-job)"; ws["F27"] = f"={S['progress']}"; ws["E27"] = "=E23*F27"
    ws["B28"] = "Final payment on completion"; ws["F28"] = "=1-F26-F27"; ws["E28"] = "=E23-E26-E27"
    for r in (26, 27, 28):
        ws.cell(row=r, column=5).number_format = MONEY; ws.cell(row=r, column=6).number_format = PCT; style_range(ws, r, r, 2, 6)
    header(ws, 30, ["#", "Detail", "", "Qty", "Unit", "Line price"]); ws.merge_cells("B30:C30")
    PL = rr("Estimate", "Proposal Line #")
    for i in range(1, N_PROP + 1):
        r = 30 + i
        m = f"MATCH({i},{PL},0)"
        ws.cell(row=r, column=1, value=f'=IF(ISNUMBER({m}),{i},"")').border = BORDER
        fx(ws, r, 2, f'=IFERROR(IF(INDEX({rr("Estimate","Description")},{m})="",INDEX({rr("Estimate","Library Item")},{m}),INDEX({rr("Estimate","Description")},{m})),"")'); ws.cell(row=r, column=3).border = BORDER
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        fx(ws, r, 4, f'=IFERROR(INDEX({rr("Estimate","Qty")},{m}),"")')
        fx(ws, r, 5, f'=IFERROR(INDEX({rr("Estimate","Unit Used")},{m}),"")')
        fx(ws, r, 6, f'=IFERROR(INDEX({EPRICE},{m}),"")', MONEY)
    tr = 31 + N_PROP
    ws.cell(row=tr, column=2, value="Lines shown"); ws.cell(row=tr, column=4, value=f'=COUNT(A31:A{tr-1})')
    ws.cell(row=tr, column=3, value=f'=IF(COUNTIF({PL},">0")>{N_PROP},"+ more lines not shown","")').font = NOTE
    fx(ws, tr, 6, f"=SUM(F31:F{tr-1})", MONEY, True); style_range(ws, tr, tr, 2, 6, fill=TOTAL_FILL, font=BOLD)
    ws.cell(row=tr + 2, column=2, value="Terms").font = Font(bold=True, color=TEAL)
    x = ws.cell(row=tr + 3, column=2, value=f"={S['proposal_terms']}"); x.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=tr + 3, start_column=2, end_row=tr + 3, end_column=6); ws.row_dimensions[tr + 3].height = 60
    ws.cell(row=tr + 4, column=2, value="Exclusions & assumptions").font = Font(bold=True, color=TEAL)
    x = ws.cell(row=tr + 5, column=2, value=f"={S['proposal_exclusions']}"); x.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=tr + 5, start_column=2, end_row=tr + 5, end_column=6); ws.row_dimensions[tr + 5].height = 48
    ws.cell(row=tr + 7, column=2, value="Accepted by (client): ______________________________    Date: ____________")
    ws.cell(row=tr + 8, column=2, value="For the contractor: _________________________________    Date: ____________")
    # Internal block (outside print area)
    ws["H4"] = "INTERNAL — not printed"; ws["H4"].font = Font(bold=True, color="C00000")
    header(ws, 13, ["Category", "Est. Cost", "Markup $"], start_col=8)
    for i, (cat, _) in enumerate(CATS, start=14):
        ws.cell(row=i, column=8, value=cat).border = BORDER
        fx(ws, i, 9, f'=SUMIFS({ECOST},{EJ},$I$2,{ECAT},B{i})', MONEY)
        fx(ws, i, 10, f'=E{i}-I{i}', MONEY)
    ws["H20"] = "Total cost"; ws["I20"] = "=SUM(I14:I19)"; ws["I20"].number_format = MONEY
    ws["H21"] = "Gross margin on subtotal"; ws["I21"] = "=IF(E20=0,0,(E20-I20)/E20)"; ws["I21"].number_format = PCT
    ws["H22"] = "Gross margin on total"; ws["I22"] = "=IF(E23=0,0,(E23-I20)/E23)"; ws["I22"].number_format = PCT
    ws["H23"] = "Estimate lines for this job"; ws["I23"] = f'=COUNTIF({EJ},$I$2)'
    for r in range(20, 24): ws.cell(row=r, column=8).font = BOLD
    print_setup(ws, landscape=False, area=f"A1:F{tr+8}")
    return ws
